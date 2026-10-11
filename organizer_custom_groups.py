import os
import re
import sys
import json
import time
import shutil
import tarfile
import tempfile
import threading
import subprocess
from pathlib import Path
from datetime import datetime
import tkinter as tk
from tkinter import ttk, filedialog, messagebox
import requests

from pypdf import PdfReader
from docx import Document

try:
    import openpyxl
except ImportError:
    openpyxl = None

try:
    from pptx import Presentation
except ImportError:
    Presentation = None

OLLAMA_HOST = "http://localhost:11434"
OLLAMA_GENERATE_URL = f"{OLLAMA_HOST}/api/generate"
OLLAMA_TAGS_URL = f"{OLLAMA_HOST}/api/tags"
OLLAMA_PULL_URL = f"{OLLAMA_HOST}/api/pull"
DEFAULT_MODEL = "llama3.2:3b"
MANIFEST_FILENAME = ".custom_organizer_history.json"

DEFAULT_CUSTOM_GROUPS = [
    "Work & Projects",
    "Financial & Tax",
    "Legal & Official",
    "Identity & Personal",
    "Resumes & Careers",
    "Media & Design",
    "Archive & Misc"
]

CONTENT_INSPECTABLE_EXTS = {
    ".pdf", ".docx", ".txt", ".md", ".csv",
    ".xlsx", ".pptx"
}

EXCLUDED_EXTENSIONS = {".crdownload", ".download", ".part", ".tmp"}
EXCLUDED_FILENAMES = {".ds_store", "thumbs.db", MANIFEST_FILENAME.lower()}


# ==============================================================================
# Automated Ollama Provisioning & Lifecycle Engine
# ==============================================================================
class OllamaManager:
    """Detects, installs, starts, and models provision Ollama on macOS without manual setup."""

    @staticmethod
    def is_api_responsive(timeout: float = 1.5) -> bool:
        try:
            r = requests.get(OLLAMA_TAGS_URL, timeout=timeout)
            return r.status_code == 200
        except Exception:
            return False

    @staticmethod
    def find_local_binary() -> str | None:
        # Standard system paths where ollama binary typically resides
        candidates = [
            shutil.which("ollama"),
            "/usr/local/bin/ollama",
            "/opt/homebrew/bin/ollama",
            str(Path.home() / ".ollama/bin/ollama"),
            "/Applications/Ollama.app/Contents/Resources/ollama"
        ]
        for c in candidates:
            if c and os.path.exists(c) and os.access(c, os.X_OK):
                return c
        return None

    @classmethod
    def start_service(cls, log_callback) -> bool:
        if cls.is_api_responsive():
            return True

        log_callback("Ollama is not responding. Checking for installed service...")

        # 1. Try launching via macOS Application bundle if installed
        if Path("/Applications/Ollama.app").exists():
            log_callback("Launching /Applications/Ollama.app in background...")
            subprocess.Popen(["open", "-a", "Ollama", "--background"])
            for _ in range(12):
                time.sleep(1)
                if cls.is_api_responsive():
                    log_callback("Ollama started successfully via application bundle.")
                    return True

        # 2. Try launching standalone CLI binary via subprocess
        binary = cls.find_local_binary()
        if binary:
            log_callback(f"Launching Ollama daemon from {binary}...")
            # Run in detached background process
            subprocess.Popen(
                [binary, "serve"],
                stdout=subprocess.DEVNULL,
                stderr=subprocess.DEVNULL,
                start_new_session=True
            )
            for _ in range(12):
                time.sleep(1)
                if cls.is_api_responsive():
                    log_callback("Ollama daemon started successfully.")
                    return True

        return False

    @classmethod
    def install_ollama_macos(cls, log_callback) -> bool:
        """Downloads and unpacks the official macOS Universal CLI archive into local user bin."""
        log_callback("Ollama is missing from this system. Initiating automatic download...")
        download_url = "https://ollama.com/download/ollama-darwin.zip"
        user_bin = Path.home() / ".local/bin"
        user_bin.mkdir(parents=True, exist_ok=True)
        dest_binary = user_bin / "ollama"

        try:
            with tempfile.TemporaryDirectory() as tmpdir:
                zip_path = Path(tmpdir) / "ollama-darwin.zip"
                log_callback("Downloading official Ollama package (~60MB)...")
                with requests.get(download_url, stream=True, timeout=60) as r:
                    r.raise_for_status()
                    with open(zip_path, "wb") as f:
                        for chunk in r.iter_content(chunk_size=1024 * 1024):
                            if chunk:
                                f.write(chunk)

                log_callback("Unpacking Ollama binary...")
                shutil.unpack_archive(str(zip_path), tmpdir)

                # Locate unpacked binary
                extracted_bin = None
                for root, _, files in os.walk(tmpdir):
                    if "ollama" in files:
                        p = Path(root) / "ollama"
                        if p.is_file():
                            extracted_bin = p
                            break

                if not extracted_bin:
                    log_callback("Failed to locate unpacked ollama binary.")
                    return False

                shutil.copy2(str(extracted_bin), str(dest_binary))
                dest_binary.chmod(0o755)

            log_callback(f"Installed Ollama to {dest_binary}")
            # Launch newly downloaded binary
            subprocess.Popen(
                [str(dest_binary), "serve"],
                stdout=subprocess.DEVNULL,
                stderr=subprocess.DEVNULL,
                start_new_session=True
            )
            for _ in range(15):
                time.sleep(1)
                if cls.is_api_responsive():
                    log_callback("Newly installed Ollama daemon is active.")
                    return True
        except Exception as e:
            log_callback(f"Failed to auto-install Ollama: {e}")

        return False

    @classmethod
    def is_model_available(cls, model_name: str) -> bool:
        try:
            r = requests.get(OLLAMA_TAGS_URL, timeout=3)
            if r.status_code == 200:
                tags = [m.get("name", "").split(":")[0] for m in r.json().get("models", [])]
                target_base = model_name.split(":")[0]
                return target_base in tags
        except Exception:
            pass
        return False

    @classmethod
    def pull_model(cls, model_name: str, log_callback, progress_callback) -> bool:
        log_callback(f"Model '{model_name}' not found locally. Initiating automated download...")
        try:
            payload = {"name": model_name, "stream": True}
            with requests.post(OLLAMA_PULL_URL, json=payload, stream=True, timeout=600) as resp:
                resp.raise_for_status()
                for line in resp.iter_lines():
                    if line:
                        chunk = json.loads(line.decode("utf-8"))
                        status = chunk.get("status", "")
                        total = chunk.get("total", 0)
                        completed = chunk.get("completed", 0)

                        if total > 0:
                            pct = int((completed / total) * 100)
                            progress_callback(f"Downloading model: {pct}% ({status})", pct)
                        else:
                            progress_callback(f"Model status: {status}", 0)
            log_callback(f"Model '{model_name}' successfully provisioned and ready.")
            return True
        except Exception as e:
            log_callback(f"Error pulling model '{model_name}': {e}")
            return False

    @classmethod
    def ensure_ready(cls, model_name: str, log_callback, progress_callback) -> bool:
        """Single comprehensive orchestrator."""
        # 1. Check or start service
        if not cls.is_api_responsive():
            if not cls.start_service(log_callback):
                if not cls.install_ollama_macos(log_callback):
                    return False

        # 2. Check or pull model
        if not cls.is_model_available(model_name):
            if not cls.pull_model(model_name, log_callback, progress_callback):
                return False

        return True


# ==============================================================================
# Document Parsing Layer
# ==============================================================================
def extract_text_pdf(filepath: Path, max_chars: int = 2500) -> str:
    try:
        reader = PdfReader(str(filepath))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception:
                return ""
        extracted = ""
        for page in reader.pages[:2]:
            extracted += page.extract_text() or ""
            if len(extracted) >= max_chars:
                break
        return extracted[:max_chars].strip()
    except Exception:
        return ""

def extract_text_docx(filepath: Path, max_chars: int = 2500) -> str:
    try:
        doc = Document(str(filepath))
        text = "\n".join([p.text for p in doc.paragraphs if p.text.strip()])
        return text[:max_chars].strip()
    except Exception:
        return ""

def extract_text_plain(filepath: Path, max_chars: int = 2500) -> str:
    try:
        with open(filepath, "r", encoding="utf-8", errors="ignore") as f:
            return f.read(max_chars).strip()
    except Exception:
        return ""

def extract_text_xlsx(filepath: Path, max_chars: int = 2000) -> str:
    if not openpyxl:
        return ""
    try:
        wb = openpyxl.load_workbook(str(filepath), read_only=True, data_only=True)
        sheet = wb.active
        rows = []
        for row in sheet.iter_rows(max_row=10, values_only=True):
            cleaned = [str(c) for c in row if c is not None]
            if cleaned:
                rows.append(", ".join(cleaned))
        return "\n".join(rows)[:max_chars].strip()
    except Exception:
        return ""

def extract_text_pptx(filepath: Path, max_chars: int = 2000) -> str:
    if not Presentation:
        return ""
    try:
        prs = Presentation(str(filepath))
        runs = []
        for slide in prs.slides[:3]:
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    runs.append(shape.text.strip())
        return "\n".join(runs)[:max_chars].strip()
    except Exception:
        return ""

def inspect_file_content(filepath: Path) -> str:
    ext = filepath.suffix.lower()
    if ext == ".pdf":
        return extract_text_pdf(filepath)
    elif ext == ".docx":
        return extract_text_docx(filepath)
    elif ext in {".txt", ".md", ".csv"}:
        return extract_text_plain(filepath)
    elif ext == ".xlsx":
        return extract_text_xlsx(filepath)
    elif ext == ".pptx":
        return extract_text_pptx(filepath)
    return ""


# ==============================================================================
# Model Query & Custom Dynamic Grouping
# ==============================================================================
def query_ollama_json(prompt: str, model: str = DEFAULT_MODEL) -> dict:
    payload = {
        "model": model,
        "prompt": prompt,
        "format": "json",
        "stream": False,
        "options": {"temperature": 0.1},
    }
    try:
        res = requests.post(OLLAMA_GENERATE_URL, json=payload, timeout=25)
        res.raise_for_status()
        return json.loads(res.json().get("response", "{}"))
    except Exception as e:
        return {"error": str(e)}

def classify_to_custom_group(filename: str, excerpt: str, custom_groups: list[str]) -> tuple[str, str]:
    groups_str = ", ".join([f'"{g}"' for g in custom_groups])
    prompt = f"""You are an archiving assistant.
Allowed Parent Groups: [{groups_str}]

Instructions:
1. Assign this document to the single best matching Parent Group from the allowed list.
2. Generate a concise, logical Subfolder name (1 to 3 words) matching the specific topic or functional domain.
   Examples:
   - For an invoice: Subfolder = "Invoices & Receipts"
   - For a tech resume: Subfolder = "Software & Engineering"
   - For a bank statement: Subfolder = "Bank Statements"
   - If no specific subfolder makes sense, use "General"

Filename: "{filename}"
Excerpt:
{excerpt if excerpt else "[No text content extracted. Determine based on filename alone]"}

Respond ONLY with this JSON schema:
{{
  "parent_group": "<Exact name from Allowed Parent Groups>",
  "subfolder": "<Concise Subfolder Name>",
  "reasoning": "<Short one-sentence justification>"
}}
"""
    data = query_ollama_json(prompt)
    parent = data.get("parent_group", "").strip()
    subfolder = data.get("subfolder", "General").strip()
    reason = data.get("reasoning", "Classified into custom group")

    parent = re.sub(r'[\\/*?:"<>|]', "", parent).strip()
    subfolder = re.sub(r'[\\/*?:"<>|]', "", subfolder).strip()

    selected_parent = None
    for g in custom_groups:
        if g.lower() == parent.lower():
            selected_parent = g
            break
    if not selected_parent:
        selected_parent = custom_groups[-1] if custom_groups else "Archive & Misc"

    if not subfolder:
        subfolder = "General"

    full_target_folder = f"{selected_parent}/{subfolder}"
    return full_target_folder, reason

def check_filename_sufficiency(filename: str) -> bool:
    generic_patterns = [r"^scan", r"^doc\d*", r"^img_\d+", r"^untitled", r"^download", r"^test", r"^\d+$"]
    stem = Path(filename).stem.lower()
    for pat in generic_patterns:
        if re.search(pat, stem):
            return False
    return len(stem) > 5


# ==============================================================================
# File Safety & History Engine
# ==============================================================================
def resolve_collision(target_dir: Path, filename: str) -> Path:
    dest = target_dir / filename
    if not dest.exists():
        return dest
    stem, suffix = Path(filename).stem, Path(filename).suffix
    c = 1
    while True:
        candidate = target_dir / f"{stem}_{c}{suffix}"
        if not candidate.exists():
            return candidate
        c += 1

def load_manifest(target_dir: Path) -> list:
    p = target_dir / MANIFEST_FILENAME
    if not p.exists():
        return []
    try:
        with open(p, "r", encoding="utf-8") as f:
            return json.load(f)
    except Exception:
        return []

def save_manifest(target_dir: Path, data: list):
    p = target_dir / MANIFEST_FILENAME
    with open(p, "w", encoding="utf-8") as f:
        json.dump(data, f, indent=2, ensure_ascii=False)


# ==============================================================================
# Modern High-Contrast GUI (Custom Groups Edition)
# ==============================================================================
class CustomGroupOrganizerApp(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("Virtual Bookkeeper AI — Zero Config Edition")
        self.geometry("1160x820")
        self.minsize(1020, 720)

        self.bg_root = "#0F1117"
        self.bg_card = "#1A1D27"
        self.bg_card_alt = "#242838"
        self.border_color = "#363C4E"

        self.fg_primary = "#FFFFFF"
        self.fg_secondary = "#CBD5E1"
        self.fg_muted = "#94A3B8"

        self.color_indigo = "#4F46E5"
        self.color_emerald = "#059669"
        self.color_gray_btn = "#334155"

        self.configure(bg=self.bg_root)
        self.selected_path = tk.StringVar(value=str(Path.home() / "Downloads"))
        self.scan_recursive = tk.BooleanVar(value=False)
        self.custom_groups = list(DEFAULT_CUSTOM_GROUPS)
        self.staged_items = []

        self._configure_styles()
        self._build_ui()

        # Run background verification of Ollama on application start
        threading.Thread(target=self._initial_engine_check, daemon=True).start()

    def _configure_styles(self):
        style = ttk.Style(self)
        style.theme_use("clam")

        style.configure(
            "Custom.Treeview",
            background=self.bg_card,
            foreground=self.fg_primary,
            fieldbackground=self.bg_card,
            rowheight=32,
            font=("SF Pro Text", 10),
            borderwidth=0
        )
        style.configure(
            "Custom.Treeview.Heading",
            background="#252A3A",
            foreground="#F8FAFC",
            font=("SF Pro Text", 10, "bold"),
            relief="flat",
            padding=(8, 8)
        )
        style.map(
            "Custom.Treeview",
            background=[("selected", "#4338CA")],
            foreground=[("selected", "#FFFFFF")]
        )

        style.configure(
            "Custom.Horizontal.TProgressbar",
            troughcolor="#1A1D27",
            background="#6366F1",
            thickness=6
        )

    def _build_ui(self):
        # Header
        header = tk.Frame(self, bg=self.bg_root, padx=24, pady=14)
        header.pack(fill=tk.X)

        title_lbl = tk.Label(
            header,
            text="✨ Virtual Bookkeeper AI",
            font=("SF Pro Display", 20, "bold"),
            fg=self.fg_primary,
            bg=self.bg_root
        )
        title_lbl.pack(side=tk.LEFT)

        subtitle_lbl = tk.Label(
            header,
            text="Automated Engine Setup • Custom Groups • Zero Configuration",
            font=("SF Pro Text", 10),
            fg=self.fg_muted,
            bg=self.bg_root
        )
        subtitle_lbl.pack(side=tk.LEFT, padx=(14, 0), pady=(6, 0))

        # 1. Target Directory Card
        dir_card = tk.Frame(self, bg=self.bg_card, highlightbackground=self.border_color, highlightthickness=1, padx=16, pady=10)
        dir_card.pack(fill=tk.X, padx=24, pady=(0, 8))

        lbl_dir = tk.Label(dir_card, text="TARGET DIRECTORY", font=("SF Pro Text", 9, "bold"), fg=self.fg_muted, bg=self.bg_card)
        lbl_dir.pack(anchor=tk.W, pady=(0, 4))

        input_row = tk.Frame(dir_card, bg=self.bg_card)
        input_row.pack(fill=tk.X)

        self.path_entry = tk.Entry(
            input_row, textvariable=self.selected_path, font=("SF Pro Text", 11),
            bg="#242838", fg=self.fg_primary, insertbackground=self.fg_primary,
            relief="flat", highlightbackground=self.border_color, highlightthickness=1
        )
        self.path_entry.pack(side=tk.LEFT, fill=tk.X, expand=True, ipady=6, padx=(0, 10))

        browse_btn = tk.Label(
            input_row, text="  Browse...  ", font=("SF Pro Text", 10, "bold"),
            bg=self.color_gray_btn, fg=self.fg_primary, relief="flat", padx=12, pady=6, cursor="pointinghand"
        )
        browse_btn.pack(side=tk.LEFT, padx=(0, 12))
        browse_btn.bind("<Button-1>", lambda e: self._browse_folder())

        check_sub = tk.Checkbutton(
            input_row, text="Include Subfolders", variable=self.scan_recursive,
            font=("SF Pro Text", 10, "bold"), fg=self.fg_secondary, bg=self.bg_card,
            selectcolor="#242838", activebackground=self.bg_card, activeforeground=self.fg_primary,
            cursor="pointinghand"
        )
        check_sub.pack(side=tk.LEFT)

        # 2. Custom Category Management Card
        cat_card = tk.Frame(self, bg=self.bg_card, highlightbackground=self.border_color, highlightthickness=1, padx=16, pady=10)
        cat_card.pack(fill=tk.X, padx=24, pady=(0, 8))

        lbl_cat = tk.Label(cat_card, text="CUSTOM PARENT GROUPS", font=("SF Pro Text", 9, "bold"), fg=self.fg_muted, bg=self.bg_card)
        lbl_cat.pack(anchor=tk.W, pady=(0, 4))

        cat_row = tk.Frame(cat_card, bg=self.bg_card)
        cat_row.pack(fill=tk.X)

        lbl_existing = tk.Label(cat_row, text="Active Groups:", font=("SF Pro Text", 10), fg=self.fg_secondary, bg=self.bg_card)
        lbl_existing.pack(side=tk.LEFT, padx=(0, 6))

        self.group_combo = ttk.Combobox(cat_row, values=self.custom_groups, state="readonly", width=22)
        if self.custom_groups:
            self.group_combo.current(0)
        self.group_combo.pack(side=tk.LEFT, padx=(0, 10))

        remove_btn = tk.Label(
            cat_row, text=" Remove Group ", font=("SF Pro Text", 9, "bold"),
            bg="#451A23", fg="#FCA5A5", padx=8, pady=4, cursor="pointinghand"
        )
        remove_btn.pack(side=tk.LEFT, padx=(0, 20))
        remove_btn.bind("<Button-1>", lambda e: self._remove_group())

        lbl_new = tk.Label(cat_row, text="Add New Group:", font=("SF Pro Text", 10), fg=self.fg_secondary, bg=self.bg_card)
        lbl_new.pack(side=tk.LEFT, padx=(0, 6))

        self.new_group_entry = tk.Entry(
            cat_row, font=("SF Pro Text", 10), bg="#242838", fg=self.fg_primary,
            insertbackground=self.fg_primary, relief="flat", highlightbackground=self.border_color, highlightthickness=1, width=20
        )
        self.new_group_entry.pack(side=tk.LEFT, ipady=4, padx=(0, 8))
        self.new_group_entry.bind("<Return>", lambda e: self._add_group())

        add_btn = tk.Label(
            cat_row, text=" + Add Group ", font=("SF Pro Text", 9, "bold"),
            bg="#1E293B", fg="#38BDF8", padx=10, pady=4, cursor="pointinghand"
        )
        add_btn.pack(side=tk.LEFT)
        add_btn.bind("<Button-1>", lambda e: self._add_group())

        # 3. Action Bar
        actions_bar = tk.Frame(self, bg=self.bg_root, padx=24, pady=4)
        actions_bar.pack(fill=tk.X)

        self.btn_step1 = tk.Label(
            actions_bar, text="Step 1: Analyze & Plan", font=("SF Pro Text", 11, "bold"),
            bg=self.color_indigo, fg="#FFFFFF", relief="flat", padx=18, pady=8, cursor="pointinghand"
        )
        self.btn_step1.pack(side=tk.LEFT, padx=(0, 12))
        self.btn_step1.bind("<Button-1>", lambda e: self._start_step1_thread())

        self.btn_step2 = tk.Label(
            actions_bar, text="Step 2: Commit & Organize", font=("SF Pro Text", 11, "bold"),
            bg="#272B38", fg="#64748B", relief="flat", padx=18, pady=8, cursor="arrow"
        )
        self.btn_step2.pack(side=tk.LEFT)
        self.btn_step2_enabled = False

        self.undo_btn = tk.Label(
            actions_bar, text="↺ Undo Last Run", font=("SF Pro Text", 10, "bold"),
            bg="#2E1B24", fg="#F87171", relief="flat", padx=14, pady=8, cursor="pointinghand"
        )
        self.undo_btn.pack(side=tk.RIGHT)
        self.undo_btn.bind("<Button-1>", lambda e: self._start_undo_thread())

        # 4. Staging Treeview
        table_card = tk.Frame(self, bg=self.bg_card, highlightbackground=self.border_color, highlightthickness=1)
        table_card.pack(fill=tk.BOTH, expand=True, padx=24, pady=8)

        columns = ("filename", "source_tier", "target_folder", "reasoning")
        self.tree = ttk.Treeview(table_card, columns=columns, show="headings", selectmode="browse", style="Custom.Treeview")
        self.tree.heading("filename", text="File Name")
        self.tree.heading("source_tier", text="Triage Method")
        self.tree.heading("target_folder", text="Custom Parent / Subfolder")
        self.tree.heading("reasoning", text="LLM Reasoning & Rules")

        self.tree.column("filename", width=250, anchor=tk.W)
        self.tree.column("source_tier", width=130, anchor=tk.CENTER)
        self.tree.column("target_folder", width=260, anchor=tk.W)
        self.tree.column("reasoning", width=380, anchor=tk.W)

        self.tree.tag_configure("oddrow", background=self.bg_card)
        self.tree.tag_configure("evenrow", background=self.bg_card_alt)

        tree_scroll = ttk.Scrollbar(table_card, orient=tk.VERTICAL, command=self.tree.yview)
        self.tree.configure(yscrollcommand=tree_scroll.set)
        self.tree.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        tree_scroll.pack(side=tk.RIGHT, fill=tk.Y)

        # 5. Status & Progress
        status_bar = tk.Frame(self, bg=self.bg_root, padx=24, pady=2)
        status_bar.pack(fill=tk.X)

        self.progress_bar = ttk.Progressbar(status_bar, mode="determinate", style="Custom.Horizontal.TProgressbar")
        self.progress_bar.pack(fill=tk.X, pady=(0, 4))

        self.status_lbl = tk.Label(status_bar, text="Initializing local environment...", font=("SF Pro Text", 10), fg=self.fg_secondary, bg=self.bg_root)
        self.status_lbl.pack(side=tk.LEFT)

        # 6. Live Log
        log_card = tk.Frame(self, bg=self.bg_card, highlightbackground=self.border_color, highlightthickness=1, padx=10, pady=6)
        log_card.pack(fill=tk.X, padx=24, pady=(0, 14))

        self.log_text = tk.Text(log_card, height=4, wrap=tk.WORD, font=("Menlo", 10), bg=self.bg_card, fg="#34D399", insertbackground="#34D399", relief="flat")
        log_scroll = ttk.Scrollbar(log_card, orient=tk.VERTICAL, command=self.log_text.yview)
        self.log_text.configure(yscrollcommand=log_scroll.set)
        self.log_text.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        log_scroll.pack(side=tk.RIGHT, fill=tk.Y)

    # --------------------------------------------------------------------------
    # Initial Engine Check
    # --------------------------------------------------------------------------
    def _initial_engine_check(self):
        ready = OllamaManager.ensure_ready(
            DEFAULT_MODEL,
            log_callback=self._log,
            progress_callback=self._update_progress_from_worker
        )
        if ready:
            self.status_lbl.config(text=f"Engine ready ({DEFAULT_MODEL}). Select directory and click Step 1.")
        else:
            self.status_lbl.config(text="Engine setup encountered an issue. Review logs.")

    def _update_progress_from_worker(self, message: str, percent: int):
        self.status_lbl.config(text=message)
        if percent > 0:
            self.progress_bar.config(mode="determinate", maximum=100, value=percent)

    # --------------------------------------------------------------------------
    # Custom Group Helpers
    # --------------------------------------------------------------------------
    def _add_group(self):
        new_name = self.new_group_entry.get().strip()
        new_clean = re.sub(r'[\\/*?:"<>|]', "", new_name).strip()
        if not new_clean:
            return
        if new_clean not in self.custom_groups:
            self.custom_groups.append(new_clean)
            self.group_combo["values"] = self.custom_groups
            self.group_combo.set(new_clean)
            self.new_group_entry.delete(0, tk.END)
            self._log(f"Added custom parent group: '{new_clean}'")
        else:
            messagebox.showinfo("Existing Group", f"'{new_clean}' is already in the list.")

    def _remove_group(self):
        selected = self.group_combo.get()
        if not selected:
            return
        if len(self.custom_groups) <= 1:
            messagebox.showwarning("Warning", "You must keep at least one parent group.")
            return
        self.custom_groups.remove(selected)
        self.group_combo["values"] = self.custom_groups
        self.group_combo.current(0)
        self._log(f"Removed custom parent group: '{selected}'")

    def _browse_folder(self):
        chosen = filedialog.askdirectory(initialdir=self.selected_path.get())
        if chosen:
            self.selected_path.set(chosen)
            self._log(f"Target directory set: {chosen}")

    def _log(self, msg: str):
        stamp = datetime.now().strftime("%H:%M:%S")
        self.log_text.insert(tk.END, f"[{stamp}] {msg}\n")
        self.log_text.see(tk.END)

    def _toggle_step2(self, active: bool):
        self.btn_step2_enabled = active
        if active:
            self.btn_step2.config(bg=self.color_emerald, fg="#FFFFFF", cursor="pointinghand")
            self.btn_step2.bind("<Button-1>", lambda e: self._start_step2_thread())
        else:
            self.btn_step2.config(bg="#272B38", fg="#64748B", cursor="arrow")
            self.btn_step2.unbind("<Button-1>")

    # --------------------------------------------------------------------------
    # Step 1: Threaded Analysis
    # --------------------------------------------------------------------------
    def _start_step1_thread(self):
        target_dir = Path(self.selected_path.get()).expanduser().resolve()
        if not target_dir.is_dir():
            messagebox.showerror("Error", f"Folder does not exist:\n{target_dir}")
            return
        if not self.custom_groups:
            messagebox.showerror("Error", "Please add at least one parent group.")
            return

        self._toggle_step2(False)
        self.tree.delete(*self.tree.get_children())
        self.staged_items.clear()
        threading.Thread(target=self._run_step1, args=(target_dir,), daemon=True).start()

    def _run_step1(self, target_dir: Path):
        start_time = time.time()
        self.status_lbl.config(text="Verifying Ollama and model status...")

        # Ensure Ollama daemon and model are online
        if not OllamaManager.ensure_ready(
            DEFAULT_MODEL,
            log_callback=self._log,
            progress_callback=self._update_progress_from_worker
        ):
            messagebox.showerror("Setup Failure", "Unable to establish local Ollama service or download model.")
            return

        pattern = "**/*" if self.scan_recursive.get() else "*"
        candidates = []
        for p in target_dir.glob(pattern):
            if not p.is_file():
                continue
            name_lower = p.name.lower()
            if name_lower in EXCLUDED_FILENAMES or name_lower.startswith(("~", "._", ".ds_store")):
                continue
            if p.suffix.lower() in EXCLUDED_EXTENSIONS:
                continue
            candidates.append(p)

        total = len(candidates)
        if total == 0:
            self._log("No eligible files found in selected directory.")
            self.status_lbl.config(text="Scan complete. Zero eligible files.")
            return

        self._log(f"Triaging {total} files across {len(self.custom_groups)} custom groups...")
        self.progress_bar.config(mode="determinate", maximum=total, value=0)

        for i, filepath in enumerate(candidates, start=1):
            fname = filepath.name
            ext = filepath.suffix.lower()
            self.status_lbl.config(text=f"Evaluating {i}/{total}: {fname}")

            if check_filename_sufficiency(fname):
                excerpt = ""
                detection_method = "Filename Priority"
            else:
                excerpt = inspect_file_content(filepath) if ext in CONTENT_INSPECTABLE_EXTS else ""
                detection_method = "Content Inspected" if excerpt else "Filename Fallback"

            target_folder, reason = classify_to_custom_group(fname, excerpt, self.custom_groups)

            item = {
                "source_path": str(filepath),
                "filename": fname,
                "detection_method": detection_method,
                "target_folder": target_folder,
                "reasoning": reason,
            }
            self.staged_items.append(item)

            row_tag = "evenrow" if i % 2 == 0 else "oddrow"
            self.tree.insert("", tk.END, values=(fname, detection_method, target_folder, reason), tags=(row_tag,))
            self.progress_bar.config(value=i)

        elapsed = round(time.time() - start_time, 2)
        self._log(f"Step 1 Complete: {total} files categorized in {elapsed}s.")
        self.status_lbl.config(text=f"Staged {total} files ({elapsed}s). Ready for Step 2 Commit.")
        self._toggle_step2(True)

    # --------------------------------------------------------------------------
    # Step 2: Threaded Commit
    # --------------------------------------------------------------------------
    def _start_step2_thread(self):
        if not self.staged_items or not self.btn_step2_enabled:
            return

        if not messagebox.askyesno(
            "Confirm Organization",
            f"Relocate {len(self.staged_items)} files into their custom folders & subfolders?\n\nAn undo snapshot will be saved to disk.",
        ):
            return

        self._toggle_step2(False)
        target_dir = Path(self.selected_path.get()).expanduser().resolve()
        threading.Thread(target=self._run_step2, args=(target_dir,), daemon=True).start()

    def _run_step2(self, target_dir: Path):
        start_time = time.time()
        self._log("Moving files into custom folder hierarchy...")
        total = len(self.staged_items)
        self.progress_bar.config(mode="determinate", maximum=total, value=0)

        session_id = datetime.now().isoformat()
        session_record = {
            "session_id": session_id,
            "timestamp": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            "directory": str(target_dir),
            "operations": [],
            "reversed": False,
        }

        moved_count = 0
        for i, item in enumerate(self.staged_items, start=1):
            src = Path(item["source_path"])
            rel_folder = item["target_folder"]

            if not src.exists():
                self._log(f"Skipped (missing): {src.name}")
                continue

            dest_folder = target_dir / rel_folder
            dest_folder.mkdir(parents=True, exist_ok=True)

            final_dest = resolve_collision(dest_folder, src.name)
            try:
                shutil.move(str(src), str(final_dest))
                session_record["operations"].append({
                    "original_path": str(src),
                    "moved_path": str(final_dest),
                    "category": rel_folder,
                })
                moved_count += 1
                self._log(f"Moved: '{src.name}' -> '{rel_folder}/{final_dest.name}'")
            except Exception as e:
                self._log(f"ERROR moving {src.name}: {e}")

            self.progress_bar.config(value=i)

        manifest = load_manifest(target_dir)
        manifest.append(session_record)
        save_manifest(target_dir, manifest)

        elapsed = round(time.time() - start_time, 2)
        self._log(f"Step 2 Finished: {moved_count} files organized in {elapsed}s.")
        self.status_lbl.config(text=f"Organized {moved_count} files successfully.")
        self.staged_items.clear()

    # --------------------------------------------------------------------------
    # Undo
    # --------------------------------------------------------------------------
    def _start_undo_thread(self):
        target_dir = Path(self.selected_path.get()).expanduser().resolve()
        manifest = load_manifest(target_dir)

        active = [s for s in manifest if not s.get("reversed", False)]
        if not active:
            messagebox.showinfo("Undo Unavailable", "No active organization sessions found to reverse.")
            return

        latest = active[-1]
        if not messagebox.askyesno(
            "Confirm Undo",
            f"Restore {len(latest['operations'])} files from session [{latest['timestamp']}] back to original locations?",
        ):
            return

        threading.Thread(target=self._run_undo, args=(target_dir, latest, manifest), daemon=True).start()

    def _run_undo(self, target_dir: Path, session: dict, manifest: list):
        self._log(f"Rolling back session: {session['session_id']}...")
        restored_count = 0
        ops = session["operations"]
        self.progress_bar.config(mode="determinate", maximum=len(ops), value=0)

        folders_to_prune = set()
        for i, op in enumerate(reversed(ops), start=1):
            moved = Path(op["moved_path"])
            orig = Path(op["original_path"])

            if moved.exists():
                orig.parent.mkdir(parents=True, exist_ok=True)
                final_orig = resolve_collision(orig.parent, orig.name)
                try:
                    shutil.move(str(moved), str(final_orig))
                    restored_count += 1
                    folders_to_prune.add(moved.parent)
                    self._log(f"Restored: '{moved.name}' -> original path")
                except Exception as e:
                    self._log(f"ERROR restoring {moved.name}: {e}")
            self.progress_bar.config(value=i)

        for folder in sorted(folders_to_prune, key=lambda p: len(str(p)), reverse=True):
            try:
                if folder.exists() and not any(folder.iterdir()):
                    folder.rmdir()
                    self._log(f"Pruned empty folder: '{folder.name}'")
                    if folder.parent != target_dir and not any(folder.parent.iterdir()):
                        folder.parent.rmdir()
            except Exception:
                pass

        session["reversed"] = True
        save_manifest(target_dir, manifest)

        self._log(f"Rollback complete: {restored_count} files restored.")
        self.status_lbl.config(text=f"Undo finished. {restored_count} files restored.")


if __name__ == "__main__":
    app = CustomGroupOrganizerApp()
    app.mainloop()